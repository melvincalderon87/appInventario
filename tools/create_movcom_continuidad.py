from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "Documentos" / "MovCom Continuidad Septiembre 2026.docx"

def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    props.append(fill)

def borders(table):
    tbl_pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "D9D9D9")
        b.append(e)
    tbl_pr.append(b)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    borders(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        shade(c, "1F4E78")
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs:
            r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); r.font.size = Pt(9)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for ri, values in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.size = Pt(9)
            if ri % 2 == 1: shade(cells[i], "F3F7FA")
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths): row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return t

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7); sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.75); sec.right_margin = Inches(0.75)
styles = doc.styles
styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(10)
styles["Title"].font.name = "Aptos Display"; styles["Title"].font.size = Pt(22); styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
for name in ("Heading 1", "Heading 2"):
    styles[name].font.name = "Aptos"; styles[name].font.color.rgb = RGBColor(0, 0, 0)
styles["Heading 1"].font.size = Pt(15); styles["Heading 2"].font.size = Pt(12)

p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("MovCom Continuidad Septiembre 2026")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Estado de integración, validación y acciones pendientes").italic = True
doc.add_paragraph("El MovCom ya incorpora el saldo inicial, los ingresos de etanol por unidad y los despachos que provienen de TMS. La continuidad requerida es mantener los cálculos de combustibles sin cambios y presentar el detalle de etanol en pestañas separadas por suplidor. La validación de septiembre confirma los ingresos de etanol; quedan diferencias menores de despacho que deben rastrearse por movimiento TMS.")

doc.add_heading("Alcance confirmado", level=1)
table(doc, ["Componente", "Regla aplicada", "Estado"], [
    ["Saldos iniciales", "Un registro por suplidor y producto con fecha 01 09 2026", "Cargado y comparado"],
    ["Ingresos de etanol", "Solo unidades. AH es real y AI es corregido. Fecha de cada recepción y unidad correlativa MMDD", "Cargado"],
    ["Recibos MovCom", "Para etanol se usa AI corregido, para cuadrar el resumen del Excel", "Integrado"],
    ["Despachos", "TMS usa ImporteComponente, signo, transacción 530, mapeo de producto y reglas de folio", "Integrado"],
    ["Combustibles", "MOGAS 95, MOGAS 91 y ULSD conservan las reglas existentes", "Sin modificación"],
], [1.55, 3.85, 1.3])

doc.add_heading("Validación de ingresos de etanol", level=1)
doc.add_paragraph("Los datos se tomaron de las hojas U Etanol por suplidor. AH corresponde al total real y AI al total corregido. El reporte usa AI para el recibo, igual que el resumen mensual del archivo MOVCOM.")
table(doc, ["Suplidor", "AG Boletas", "AH Real", "AI Corregido", "Registros"], [
    ["Chevron Guatemala", "526,751.486", "527,317", "518,571", "13"],
    ["Uno Fuels", "344,000", "344,330", "337,769", "4"],
    ["UNO Guatemala", "593,500", "594,170", "582,873", "11"],
], [1.8, 1.15, 1.15, 1.3, 0.8])

doc.add_heading("Revisión de despacho TMS", level=1)
doc.add_paragraph("El cálculo vigente de etanol en MovCom parte de TMS y no de los ingresos por unidad. La diferencia mostrada es pequeña, pero debe revisarse antes de considerar el cierre completamente conciliado.")
table(doc, ["Suplidor", "Excel galones", "TMS MovCom galones", "Diferencia"], [
    ["UNO Guatemala", "273,729", "273,811", "+82"],
    ["Uno Fuels", "84,317", "84,411", "+94"],
    ["Chevron Guatemala", "275,070", "275,119", "+49"],
], [2.2, 1.4, 1.8, 1.0])

doc.add_heading("Nueva presentación de etanol", level=1)
doc.add_paragraph("Se agrega una pestaña nueva por suplidor en la descarga de MovCom: Etanol UNO Guatemala, Etanol Uno Fuels y Etanol Chevron Guatemala. Cada pestaña muestra la fecha, la unidad correlativa, AG Boletas, AH Real y AI Corregido en galones. Estas pestañas son adicionales; no cambian las pestañas, fórmulas ni cálculos existentes de combustibles.")

doc.add_heading("Pendientes para cierre correcto", level=1)
bullet(doc, "Rastrear los 82, 94 y 49 galones de diferencia por folio, orden, hora de carga y línea TMS antes de cerrar septiembre.")
bullet(doc, "Registrar AG Boletas también en las futuras entradas de etanol. Los datos históricos de septiembre ya se cargaron para que las nuevas pestañas repliquen el Excel.")
bullet(doc, "Validar los saldos físicos de etanol y tubería contra el cierre final si se necesita calcular la ganancia o pérdida de etanol.")
bullet(doc, "Mantener el correlativo MMDD mientras no se disponga de la identificación real de cada unidad; sustituirlo cuando se reciba el identificador oficial.")

doc.add_heading("Criterio de aceptación", level=1)
doc.add_paragraph("El cierre queda listo para operación cuando los recibos AI del sistema coincidan con el Excel, las nuevas pestañas de etanol muestren los ingresos por unidad y las diferencias de despacho TMS estén justificadas o corregidas por su folio de origen.")

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
